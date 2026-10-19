# -*- coding: utf-8 -*-
"""生成罗常青的实习简历（DOCX）。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

EA = "微软雅黑"
LATIN = "Segoe UI"
ACCENT = RGBColor(0x1B, 0x3A, 0x5C)
BODY = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x63, 0x63, 0x63)
RULE = "9DB4D2"

CONTENT_W = 17.2  # cm
BODY_SIZE = 9.8
LINE = 1.03
HEAD_BEFORE = 6
HEAD_AFTER = 2.5

_PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:jc", "w:rPr", "w:sectPr", "w:pPrChange",
)


def style_run(r, size=BODY_SIZE, bold=False, color=BODY, italic=False):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    rFonts = r._element.get_or_add_rPr().get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), LATIN)
    rFonts.set(qn("w:hAnsi"), LATIN)
    rFonts.set(qn("w:eastAsia"), EA)
    rFonts.set(qn("w:cs"), LATIN)
    return r


def para(doc, before=0, after=0, line=LINE, left=0.0, hanging=0.0, right_tab=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if left:
        pf.left_indent = Cm(left)
    if hanging:
        pf.first_line_indent = Cm(-hanging)
    if right_tab:
        pf.tab_stops.add_tab_stop(Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
    return p


def put(p, text, size=BODY_SIZE, bold=False, color=BODY, italic=False):
    return style_run(p.add_run(text), size, bold, color, italic)


def bottom_rule(p, color=RULE, size=8, space=3):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.insert_element_before(pBdr, *_PBDR_SUCCESSORS)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)


def heading(doc, text):
    p = para(doc, before=HEAD_BEFORE, after=HEAD_AFTER)
    put(p, text, 11.5, True, ACCENT)
    bottom_rule(p)
    return p


def bullet(doc, label, text, after=0.8):
    p = para(doc, after=after, left=0.45, hanging=0.45)
    put(p, "• ", BODY_SIZE, False, ACCENT)
    if label:
        put(p, label, BODY_SIZE, True, BODY)
    if text:
        put(p, text, BODY_SIZE, False, BODY)
    return p


def entry(doc, left_text, right_text, before=2):
    p = para(doc, before=before, after=0, right_tab=True)
    put(p, left_text, 10.5, True, BODY)
    put(p, "\t" + right_text, 9.5, False, GRAY)
    return p


def build(path):
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = LATIN
    st.font.size = Pt(BODY_SIZE)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EA)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = LINE

    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(0.85)
    sec.left_margin = Cm(1.9)
    sec.right_margin = Cm(1.9)

    # ---------- 页眉 ----------
    p = para(doc, after=2, right_tab=True)
    put(p, "罗常青", 22, True, ACCENT)
    put(p, "\t男 · 2003.10 · 汉族 · 共青团员 · 重庆万州", 9.5, False, GRAY)

    p = para(doc, after=1)
    put(p, "求职意向：前端开发实习生（前端 / 前端 + Agent 开发方向）", 10.5, True, ACCENT)

    p = para(doc, after=0)
    put(p, "【手机号】 · 【邮箱】 · GitHub：【链接】 · 可实习时间：【待填】", 9.5, False, GRAY)
    bottom_rule(p, space=4)

    # ---------- 教育背景 ----------
    heading(doc, "教育背景")
    entry(doc, "重庆邮电大学　计算机技术（专业硕士，在读）", "2025.09 – 2028.06")
    bullet(doc, "奖学金：",
           "研究生学业奖学金一等奖（研一学年）、二等奖（研0入学）；预计 2028.06 毕业",
           after=3)

    entry(doc, "重庆邮电大学　计算机科学与技术学院（本科）", "2021.09 – 2025.06")
    bullet(doc, "奖学金：", "本科学业奖学金三等奖")

    # ---------- 专业技能 ----------
    heading(doc, "专业技能")
    bullet(doc, "前端基础：",
           "HTML / CSS / JavaScript（ES6+、闭包、原型、事件循环）、TypeScript；"
           "浏览器渲染流程与性能优化、响应式布局与移动端适配")
    bullet(doc, "框架与工程化：",
           "React（Hooks、状态管理、渲染优化）；Next.js 15（App Router、Server / Client 组件、"
           "Route Handlers）；Tailwind CSS、shadcn/ui；Zustand、React Context；Vitest")
    bullet(doc, "富文本与编辑器：",
           "BlockNote（底层 ProseMirror）JSON 数据驱动模型、onChange 增量保存与大文档性能边界处理")
    bullet(doc, "Agent 开发：",
           "Vercel AI SDK、DeepSeek 模型接入；Tool / Function Calling、Agent Loop 与停止条件；"
           "Prompt 工程；SSE 流式输出；上下文窗口管理与摘要压缩；了解 MCP / RAG")
    bullet(doc, "后端与数据（了解层面）：",
           "Supabase（Auth / Postgres / RLS）、PostgreSQL 单表自引用树建模")

    # ---------- 项目经历 ----------
    heading(doc, "项目经历")
    entry(doc, "MyLotion｜全栈 AI 笔记应用（个人项目）", "2026.09 – 至今")
    p = para(doc, after=2.5)
    put(p, "技术栈：Next.js 15 · React · TypeScript · BlockNote · Tailwind CSS · Supabase · "
           "Vercel AI SDK · DeepSeek · SSE", 9.5, False, GRAY)

    bullet(doc, "",
           "独立完成侧边栏文档树、BlockNote 编辑器、AI 对话面板三大模块及后端 API，"
           "打通「自然语言 → 模型调工具 → 数据库改写 → 前端实时刷新」完整链路。")
    bullet(doc, "无限层级文档树：",
           "PostgreSQL 单表 parentDocument 自引用建模；侧边栏一次拉取扁平数组，"
           "前端两遍 Map 扫描 O(n) 还原成树，替代递归 filter 的 O(n²)。")
    bullet(doc, "子树批量归档：",
           "把「逐层查库」的递归（N+1）改为内存 BFS 一次 O(n) 算出全部后代 id + 单次批量更新，"
           "百级节点归档从上百次网络往返降到 2 次。")
    bullet(doc, "Agent 编排：",
           "streamText + DeepSeek 驱动「思考 → 调工具 → 看结果」循环，定义搜索 / 读取 / 新建 / 更新 / "
           "重命名 / 归档 / 删除 7 个 Tool，用 stopWhen 限制最大步数；工具由 withToolLogging 统一包装，"
           "异常转成可读错误回传模型重试，并设 30s 超时。")
    bullet(doc, "流式通信：",
           "工具在服务端执行、前端不可见，副作用由注入的 onEvent 回调写入事件队列，SSE 每 200ms 编码为 "
           "data: <json> 推给前端，前端按 type 分发（文本 / 进度 / 刷新 / 确认弹窗）；选 SSE 而非 "
           "WebSocket：单向推送足够、基于 HTTP 免协议升级、自动重连。")
    bullet(doc, "安全与幂等：",
           "删除类工具不直接落库，只上报 confirm_delete，用户确认后才由前端执行（AI 只能发起、"
           "用户才能执行）；requestId + 唯一约束实现消息幂等，重复提交返回 409；跨用户数据由 RLS 兜底。")
    bullet(doc, "上下文与工程质量：",
           "「系统提示词 + 历史摘要 + 最近 100 条原文」的确定性滑动窗口，窗口外消息由模型重写为摘要以控 "
           "token；文档详情读取做并发请求去重 + LRU（上限 200）缓存；Vitest 覆盖纯函数与算法用例。")

    # ---------- 校园经历 ----------
    heading(doc, "校园经历")
    bullet(doc, "学院学生会 体育部负责人（研一）：",
           "统筹院级体育活动的策划、人员安排与现场执行，对接学院与各班级。")
    bullet(doc, "学院足球队（本科）：",
           "长期参加训练与院际比赛，具备团队协作意识与抗压能力。")

    doc.save(path)
    print("saved:", path)


if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "罗常青-前端开发实习生-简历.docx")
    build(out)
